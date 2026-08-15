from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any, Iterable
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def _new_document(title: str) -> Document:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return document


def _bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_key_values(document: Document, items: Iterable[tuple[str, Any]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in items:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)


def build_group_roster_csv(result: dict[str, Any]) -> bytes:
    output = StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=["student_id", "group", "accuracy", "average_response_time_sec", "main_error", "slow_correct_count", "fast_wrong_count"],
    )
    writer.writeheader()
    writer.writerows(result.get("class_profile", {}).get("student_details", []))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def build_feedback_students_csv(feedback: dict[str, Any]) -> bytes:
    output = StringIO()
    fields = ["student_id", "knowledge_point", "before_accuracy", "after_accuracy", "change", "main_error", "after_group", "continue_intervention"]
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    writer.writerows(feedback.get("student_changes", []))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def build_analysis_docx(result: dict[str, Any], feedback: dict[str, Any] | None = None) -> bytes:
    profile = result["class_profile"]
    document = _new_document("班级学情分析报告")
    _add_key_values(
        document,
        [
            (
                "学科",
                result["teacher_request"].get(
                    "subject", result.get("lesson_plan", {}).get("subject", "")
                ),
            ),
            ("课题", result["teacher_request"].get("topic", "")),
            ("学生数", profile["student_count"]),
            ("答题记录", profile["record_count"]),
            ("课前正确率", f"{profile['overall_accuracy']:.1%}"),
        ],
    )
    document.add_heading("知识点掌握情况", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["知识点", "作答数", "正确率", "平均用时", "受影响学生", "主要错因"]):
        cell.text = label
    for item in profile["knowledge_points"]:
        cells = table.add_row().cells
        values = [item["name"], item["attempts"], f"{item['accuracy']:.1%}", f"{item.get('average_response_time_sec', 0)} 秒", item["affected_students"], item["main_error"]]
        for cell, value in zip(cells, values):
            cell.text = str(value)
    document.add_heading("主要错因", level=1)
    for item in profile["error_distribution"]:
        document.add_paragraph(f"{item['name']}：{item['count']} 次", style="List Bullet")
    document.add_heading("分组花名册", level=1)
    roster = document.add_table(rows=1, cols=7)
    roster.style = "Table Grid"
    labels = ["学生编号", "分组", "正确率", "平均用时", "主要错因", "正确但偏慢", "错误且过快"]
    for cell, label in zip(roster.rows[0].cells, labels):
        cell.text = label
    for item in profile.get("student_details", []):
        values = [item["student_id"], item["group"], f"{item['accuracy']:.1%}", f"{item['average_response_time_sec']} 秒", item["main_error"], item["slow_correct_count"], item["fast_wrong_count"]]
        for cell, value in zip(roster.add_row().cells, values):
            cell.text = str(value)
    time_analysis = profile.get("time_analysis", {})
    document.add_heading("作答时间分析", level=1)
    document.add_paragraph(
        f"班级平均作答时间：{time_analysis.get('overall_average_response_time_sec', 0)} 秒；"
        f"正确但耗时偏长：{time_analysis.get('slow_correct_count', 0)} 条；"
        f"错误且作答过快：{time_analysis.get('fast_wrong_count', 0)} 条。"
    )
    if feedback:
        document.add_heading("课前课后效果", level=1)
        document.add_paragraph(
            f"课后正确率 {feedback['after_overall_accuracy']:.1%}，"
            f"较课前变化 {feedback['overall_change']:+.1%}。"
        )
    document.add_paragraph(profile["diagnosis_note"])
    return _bytes(document)


def build_lesson_docx(result: dict[str, Any]) -> bytes:
    plan = result["lesson_plan"]
    document = _new_document(plan["title"])
    _add_key_values(
        document,
        [
            ("学科", plan.get("subject", "")),
            ("年级", plan["grade"]),
            ("课时", f"{plan['duration_minutes']} 分钟"),
        ],
    )
    document.add_heading("教学目标", level=1)
    for item in plan["objectives"]:
        document.add_paragraph(str(item), style="List Number")
    document.add_heading("教学策略", level=1)
    for item in plan["key_strategies"]:
        document.add_paragraph(str(item), style="List Bullet")
    if plan.get("classroom_questions"):
        document.add_heading("课堂提问", level=1)
        for item in plan["classroom_questions"]:
            document.add_paragraph(str(item), style="List Bullet")
    document.add_heading("课堂流程", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["时间", "环节", "活动"]):
        cell.text = label
    for item in plan["timeline"]:
        cells = table.add_row().cells
        cells[0].text = f"{item['minutes']} 分钟"
        cells[1].text = str(item["stage"])
        cells[2].text = str(item["activity"])
    document.add_heading("自动质量审查", level=1)
    report = result.get("quality_report", {})
    document.add_paragraph("通过" if report.get("passed") else "需要教师复核")
    for check in report.get("checks", []):
        document.add_paragraph(str(check), style="List Bullet")
    for issue in report.get("issues", []):
        document.add_paragraph(f"提示：{issue}")
    document.add_heading("教材与资料依据", level=1)
    for item in plan.get("source_evidence", []):
        page = f"第 {item['page']} 页" if item.get("page") else f"片段 {item.get('section', '')}"
        document.add_paragraph(f"{item['source']}，{page}：{item['text']}", style="List Bullet")
    document.add_paragraph(plan["teacher_boundary"])
    return _bytes(document)


def build_practice_docx(result: dict[str, Any]) -> bytes:
    materials = result["practice_sets"]
    document = _new_document("分层练习")
    document.add_paragraph(f"重点错因：{materials['focus_error']}")
    for group in ("A_基础巩固", "B_重点纠错", "C_拓展提升"):
        document.add_heading(group, level=1)
        for index, item in enumerate(materials[group], start=1):
            document.add_paragraph(f"{index}. {item['question']}")
            document.add_paragraph(f"设计目的：{item.get('purpose', '')}")
        document.add_heading(f"{group} 参考答案", level=2)
        for index, item in enumerate(materials[group], start=1):
            document.add_paragraph(f"{index}. {item['answer']}")
    return _bytes(document)


def build_exit_ticket_docx(result: dict[str, Any]) -> bytes:
    ticket = result["practice_sets"]["exit_ticket"]
    document = _new_document("随堂测及答案")
    document.add_heading("随堂测", level=1)
    document.add_paragraph(ticket["question"])
    document.add_heading("检查要点", level=1)
    for item in ticket.get("checks", []):
        document.add_paragraph(str(item), style="List Bullet")
    document.add_page_break()
    document.add_heading("参考答案", level=1)
    document.add_paragraph(ticket["answer"])
    return _bytes(document)


def build_next_lesson_docx(feedback: dict[str, Any] | None) -> bytes:
    document = _new_document("下一课调整建议")
    if not feedback:
        document.add_paragraph("尚未上传随堂测，完成课后反馈后将生成下一课调整建议。")
        return _bytes(document)
    _add_key_values(
        document,
        [
            ("课前正确率", f"{feedback['before_overall_accuracy']:.1%}"),
            ("课后正确率", f"{feedback['after_overall_accuracy']:.1%}"),
            ("整体变化", f"{feedback['overall_change']:+.1%}"),
        ],
    )
    document.add_heading("需要继续干预的学生", level=1)
    document.add_paragraph("、".join(feedback["continued_intervention_students"]) or "暂无")
    if feedback.get("student_changes"):
        document.add_heading("学生级前后变化", level=1)
        table = document.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ["学生", "知识点", "课前", "课后", "变化", "主要错因", "课后分组"]):
            cell.text = label
        for item in feedback["student_changes"]:
            def pct(value: Any) -> str:
                return "未测" if value is None else f"{float(value):.1%}"
            values = [item["student_id"], item["knowledge_point"], pct(item["before_accuracy"]), pct(item["after_accuracy"]), pct(item["change"]), item["main_error"], item["after_group"]]
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value)
    document.add_heading("下一课建议", level=1)
    for item in feedback["next_lesson_suggestions"]:
        document.add_paragraph(str(item), style="List Bullet")
    document.add_paragraph(feedback["boundary"])
    return _bytes(document)


def build_delivery_zip(result: dict[str, Any], feedback: dict[str, Any] | None = None) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("01_学情分析报告.docx", build_analysis_docx(result, feedback))
        archive.writestr("02_教师最终教案.docx", build_lesson_docx(result))
        archive.writestr("03_分层练习及答案.docx", build_practice_docx(result))
        archive.writestr("04_随堂测及答案.docx", build_exit_ticket_docx(result))
        archive.writestr("05_下一课调整建议.docx", build_next_lesson_docx(feedback))
    return buffer.getvalue()
