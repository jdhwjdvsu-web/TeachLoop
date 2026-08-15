import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import pandas as pd
from docx import Document

from teachloop.diagnosis import diagnose_class, infer_error_type, validate_records
from teachloop.feedback import compare_learning_outcomes
from teachloop.knowledge_base import LocalVectorKnowledgeBase
from teachloop.lesson import review_output
from teachloop.llm_service import (
    EducationLLMService,
    LLMConfig,
    OpenAICompatibleClient,
    validate_generated_item,
)
from teachloop.math_validation import (
    equivalent_equations,
    number_line_description,
    validate_inequality_answer,
    validate_solution,
)
from teachloop.skill_registry import orchestration_manifest, selected_skills
from teachloop.storage import TeachLoopStore
from teachloop.subject_packs import (
    demo_path,
    grades_for_subject,
    subject_names,
    topics_for_subject,
    validate_subject_item,
)
from teachloop.word_export import build_delivery_zip, build_group_roster_csv
from teachloop.workflow import run_teaching_workflow

ROOT = Path(__file__).parents[1]


def demo_records():
    return pd.read_csv(ROOT / "data" / "demo_class.csv").to_dict(orient="records")


def test_math_solution_validation():
    assert validate_solution("2*x+3=7", "2")
    assert not validate_solution("2*x+3=7", "5")
    assert equivalent_equations("2*x+3=7", "2*x=4")


def test_error_inference():
    row = {"correct": False, "student_work": "2*x+3=7 -> 2*x=7+3", "student_answer": "5"}
    assert infer_error_type(row) == "移项忘记变号"


def test_class_diagnosis():
    profile = diagnose_class(demo_records())
    assert profile["student_count"] == 6
    assert profile["record_count"] == 24
    assert profile["error_distribution"]


def test_full_workflow():
    result = run_teaching_workflow(
        demo_records(),
        {"grade": "七年级", "topic": "一元一次方程", "duration": 45},
    )
    assert not result["validation_errors"]
    assert result["quality_report"]["passed"] is True
    assert result["lesson_plan"]["duration_minutes"] == 45


def test_streamlit_smoke():
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=30)
    assert not app.exception
    app.button[0].click().run(timeout=30)
    assert not app.exception
    assert [tab.label for tab in app.tabs] == ["班级学情", "差异化教案", "分层练习", "质量审查"]


def test_feedback_closes_the_loop():
    before = diagnose_class(demo_records())
    after_records = pd.read_csv(ROOT / "data" / "demo_post_class.csv").to_dict(orient="records")
    feedback = compare_learning_outcomes(before, after_records)
    assert feedback["after_overall_accuracy"] > feedback["before_overall_accuracy"]
    assert feedback["overall_change"] == 0.25
    assert "S006" in feedback["continued_intervention_students"]


def test_local_vector_knowledge_base(tmp_path):
    knowledge_base = LocalVectorKnowledgeBase(tmp_path / "index.json")
    added = knowledge_base.add_text(
        "移项是等式两边同时加上或减去同一个数的简写。移项后需要改变符号。",
        source="测试教材",
    )
    assert added == 1
    assert knowledge_base.add_text(
        "移项是等式两边同时加上或减去同一个数的简写。移项后需要改变符号。",
        source="测试教材",
    ) == 0
    results = knowledge_base.search("移项变号 等式性质")
    assert results
    assert results[0]["source"] == "测试教材"


def test_teacher_version_storage(tmp_path):
    store = TeachLoopStore(tmp_path / "teachloop.db")
    version_id = store.save_version("测试教案", "教师已确认", {"lesson_plan": {"title": "测试教案"}})
    versions = store.list_versions()
    assert version_id == 1
    assert versions[0]["status"] == "教师已确认"


def test_inequality_validation_and_number_line():
    assert validate_inequality_answer("-2*x > 6", "x < -3")
    assert not validate_inequality_answer("-2*x > 6", "x > -3")
    number_line = number_line_description("x <= 3")
    assert number_line == {"boundary": 3.0, "closed": True, "direction": "left", "operator": "<="}


def test_inequality_workflow():
    records = pd.read_csv(ROOT / "data" / "demo_inequality.csv").to_dict(orient="records")
    result = run_teaching_workflow(
        records,
        {"grade": "八年级", "topic": "一元一次不等式", "duration": 45},
    )
    assert result["class_profile"]["topic"] == "一元一次不等式"
    assert result["practice_sets"]["exit_ticket"]["answer"] == "x < -3"
    assert validate_generated_item(
        "一元一次不等式",
        {"question": "解不等式：-4*x <= 12", "answer": "x >= -3"},
    )


def test_openai_compatible_client_json():
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path.endswith("/chat/completions")
        assert request.headers["Authorization"] == "Bearer test-key"
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"message": {"content": "```json\n{\"objectives\": [\"目标1\"]}\n```"}}
                ]
            },
        )

    client = OpenAICompatibleClient(
        LLMConfig(base_url="https://example.test/v1", model="test-model", api_key="test-key"),
        transport=httpx.MockTransport(handler),
    )
    assert client.complete_json("system", "user")["objectives"] == ["目标1"]


def test_word_delivery_package():
    result = run_teaching_workflow(
        demo_records(),
        {"grade": "七年级", "topic": "一元一次方程", "duration": 45},
    )
    package = build_delivery_zip(result)
    with ZipFile(BytesIO(package)) as archive:
        names = archive.namelist()
        assert len(names) == 5
        assert "02_教师最终教案.docx" in names
        assert all(archive.read(name).startswith(b"PK") for name in names)


def test_workflow_with_mock_llm_enrichment():
    generated = {
        "complex_step_findings": ["学生混淆了移项与等式两边同时操作。"],
        "objectives": ["模型增强目标"],
        "key_strategies": ["模型增强策略"],
        "classroom_questions": ["为什么移项后符号会改变？"],
        "timeline": [
            {"minutes": 5, "stage": "导入", "activity": "错例"},
            {"minutes": 40, "stage": "学习", "activity": "讲练结合"},
        ],
        "practice_sets": {
            "A_基础巩固": [
                {"question": "解方程：2*x=4", "answer": "x=2", "purpose": "基础"},
                {"question": "解方程：x+1=3", "answer": "x=2", "purpose": "基础"},
            ],
            "B_重点纠错": [
                {"question": "解方程：3*x=9", "answer": "x=3", "purpose": "纠错"},
                {"question": "解方程：x-2=5", "answer": "x=7", "purpose": "纠错"},
            ],
            "C_拓展提升": [
                {"question": "解方程：4*x+1=9", "answer": "x=2", "purpose": "拓展"},
                {"question": "解方程：5*x-5=10", "answer": "x=3", "purpose": "拓展"},
            ],
        },
    }

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={"choices": [{"message": {"content": json.dumps(generated, ensure_ascii=False)}}]},
        )

    client = OpenAICompatibleClient(
        LLMConfig(base_url="https://example.test/v1", model="mock-model"),
        transport=httpx.MockTransport(handler),
    )
    result = run_teaching_workflow(
        demo_records(),
        {"grade": "七年级", "topic": "一元一次方程", "duration": 45},
        llm_service=EducationLLMService(client),
    )
    assert result["llm_status"]["used"] is True
    assert result["lesson_plan"]["objectives"] == ["模型增强目标"]
    assert result["practice_sets"]["A_基础巩固"][0]["answer"] == "x=2"


def test_selected_education_skills_and_subject_catalog():
    assert len(selected_skills()) == 12
    assert subject_names() == ["数学", "语文", "英语", "物理"]
    assert "记叙文阅读：人物形象分析" in topics_for_subject("语文", "七年级")
    assert grades_for_subject("物理") == ["八年级"]
    manifest = orchestration_manifest("语文")
    assert manifest["license"] == "CC BY-SA 4.0"
    assert manifest["stages"]["diagnosis"]


def test_chinese_subject_pack_workflow():
    records = pd.read_csv(demo_path("语文")).to_dict(orient="records")
    result = run_teaching_workflow(
        records,
        {
            "subject": "语文",
            "grade": "七年级",
            "topic": "记叙文阅读：人物形象分析",
            "duration": 45,
        },
    )
    assert result["quality_report"]["passed"] is True
    assert result["lesson_plan"]["subject"] == "语文"
    assert result["practice_sets"]["validation_mode"] == "量规＋教师复核"
    assert result["capability_context"]["subject_pack"]["validator"] == "量规与教师复核"


def test_physics_subject_pack_validation_and_workflow():
    item = {
        "question": "物体运动100 m用时20 s，求速度。",
        "answer": "5 m/s",
        "purpose": "直接使用速度公式",
    }
    assert validate_subject_item("物理", "速度与运动描述", item)
    assert not validate_subject_item("物理", "速度与运动描述", {**item, "answer": "5"})
    records = pd.read_csv(demo_path("物理")).to_dict(orient="records")
    result = run_teaching_workflow(
        records,
        {"subject": "物理", "grade": "八年级", "topic": "速度与运动描述", "duration": 45},
    )
    assert result["quality_report"]["passed"] is True
    assert result["class_profile"]["error_distribution"][0]["name"] in {
        "单位换算错误",
        "平均速度直接取平均",
    }


def test_yaml_only_english_pack_drives_error_inference():
    row = {
        "correct": False,
        "student_answer": "He go to school.",
        "student_work": "he go",
        "knowledge_point": "第三人称单数",
    }
    assert infer_error_type(row, subject="英语") == "主谓一致错误"
    records = pd.read_csv(demo_path("英语")).to_dict(orient="records")
    result = run_teaching_workflow(
        records,
        {"subject": "英语", "grade": "七年级", "topic": "一般现在时：主谓一致", "duration": 45},
    )
    assert result["class_profile"]["error_distribution"][0]["name"] == "主谓一致错误"


def test_record_validation_reports_dirty_row_details():
    records = demo_records()[:2]
    records[1] = {**records[1], "correct": "maybe", "response_time_sec": "fast"}
    records.append({**records[0]})
    errors = validate_records(records)
    assert any("第 3 行" in item and "correct" in item for item in errors)
    assert any("第 3 行" in item and "response_time_sec" in item for item in errors)
    assert any("重复" in item for item in errors)


def test_response_time_changes_student_group_and_is_reported():
    records = [
        {"student_id": "FAST", "question_id": "Q1", "knowledge_point": "K", "student_answer": "1", "student_work": "", "correct": True, "response_time_sec": 20},
        {"student_id": "SLOW", "question_id": "Q1", "knowledge_point": "K", "student_answer": "1", "student_work": "", "correct": True, "response_time_sec": 200},
        {"student_id": "FAST", "question_id": "Q2", "knowledge_point": "K", "student_answer": "1", "student_work": "", "correct": True, "response_time_sec": 20},
        {"student_id": "SLOW", "question_id": "Q2", "knowledge_point": "K", "student_answer": "1", "student_work": "", "correct": True, "response_time_sec": 200},
    ]
    profile = diagnose_class(records, subject="数学")
    assert "FAST" in profile["student_groups"]["C_拓展提升"]
    assert "SLOW" in profile["student_groups"]["B_重点纠错"]
    assert profile["time_analysis"]["slow_correct_count"] == 2


def test_feedback_anchor_and_unmeasured_knowledge_point():
    before = diagnose_class(demo_records(), subject="数学", topic="一元一次方程")
    after = [
        {"student_id": "S001", "question_id": "POST1", "linked_pre_question_id": "Q1", "knowledge_point": "移项与合并同类项", "student_answer": "2", "student_work": "", "correct": True, "response_time_sec": 30},
        {"student_id": "S002", "question_id": "POST1", "linked_pre_question_id": "Q1", "knowledge_point": "移项与合并同类项", "student_answer": "2", "student_work": "", "correct": True, "response_time_sec": 35},
    ]
    feedback = compare_learning_outcomes(before, after)
    assert feedback["question_anchor_changes"][0]["mapped"] is True
    unmeasured = next(item for item in feedback["knowledge_changes"] if item["after_accuracy"] is None)
    assert unmeasured["status"] == "未测"
    assert feedback["comparison_warnings"]


def test_group_roster_is_exported_to_csv_and_word():
    result = run_teaching_workflow(
        demo_records(), {"grade": "七年级", "topic": "一元一次方程", "duration": 45}
    )
    csv_text = build_group_roster_csv(result).decode("utf-8-sig")
    assert "student_id,group" in csv_text
    package = build_delivery_zip(result)
    with ZipFile(BytesIO(package)) as archive:
        document = Document(BytesIO(archive.read("01_学情分析报告.docx")))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "分组花名册" in text
    assert "S001" in table_text


def test_quality_review_rejects_unrelated_strategy():
    result = run_teaching_workflow(
        demo_records(), {"grade": "七年级", "topic": "一元一次方程", "duration": 45}
    )
    result["lesson_plan"]["key_strategies"] = ["播放一段与本课无关的视频。"]
    report = review_output(result["lesson_plan"], result["practice_sets"])
    assert report["passed"] is False
    assert "教学策略未对准班级最高频错因" in report["issues"]


def test_bm25_knowledge_base_source_management(tmp_path):
    kb = LocalVectorKnowledgeBase(tmp_path / "kb.json")
    kb.add_text("移项变号来源于等式两边同时进行相同操作。", "数学资料")
    kb.add_text("人物形象分析需要观点、证据和解释。", "语文资料")
    assert kb.search("等式 移项变号")[0]["source"] == "数学资料"
    assert kb.rename_source("数学资料", "数学教材") == 1
    assert kb.search("移项")[0]["source"] == "数学教材"
    assert kb.delete_source("数学教材") == 1
    assert all(item["source"] != "数学教材" for item in kb.search("移项", top_k=10))


def test_llm_retries_after_rate_limit():
    calls = {"count": 0}
    def handler(request: httpx.Request) -> httpx.Response:
        calls["count"] += 1
        if calls["count"] == 1:
            return httpx.Response(429, json={"error": "rate limited"})
        return httpx.Response(200, json={"choices": [{"message": {"content": '{"ok": true}'}}], "usage": {"total_tokens": 4}})
    client = OpenAICompatibleClient(
        LLMConfig(base_url="https://example.test/v1", model="test", max_retries=2),
        transport=httpx.MockTransport(handler),
    )
    assert client.complete_json("system", "user")["ok"] is True
    assert client.last_call_metrics["attempts"] == 2


def test_connection_falls_back_to_chat_when_models_missing():
    def handler(request: httpx.Request) -> httpx.Response:
        if request.method == "GET":
            return httpx.Response(404, json={"error": "unsupported"})
        return httpx.Response(200, json={"choices": [{"message": {"content": '{"ok": true}'}}]})
    client = OpenAICompatibleClient(
        LLMConfig(base_url="https://example.test/v1", model="test"),
        transport=httpx.MockTransport(handler),
    )
    status = client.test_connection()
    assert status["ok"] is True
    assert status["method"] == "chat"


def test_feedback_history_and_schema_version(tmp_path):
    store = TeachLoopStore(tmp_path / "history.db")
    feedback_id = store.save_feedback("测试课", {"overall_change": 0.1})
    feedbacks = store.list_feedbacks()
    assert feedbacks[0]["id"] == feedback_id
    assert feedbacks[0]["payload"]["schema_version"] == 2
